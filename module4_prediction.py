# 模块4代码
import os
os.environ["AKSHARE_JS_ENGINE"] = "nodejs"
import json
import jieba
import numpy as np
import pandas as pd
from datetime import datetime, timedelta
from PyQt6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
                             QLabel, QComboBox, QTextEdit, QFrame, QFileDialog,
                             QListWidget, QScrollArea, QLineEdit)
from PyQt6.QtCore import Qt, QTimer, QThread, pyqtSignal, QPropertyAnimation
from PyQt6.QtGui import QDragEnterEvent, QDropEvent
from utils import ZHIPU_API_KEY, MODULE1_MODEL
from zhipuai import ZhipuAI
from utils import get_resource_path

try:
    import akshare as ak
    AKSHARE_AVAILABLE = True
except ImportError:
    AKSHARE_AVAILABLE = False

def load_sentiment_dict():
    dict_path = get_resource_path("resources/sentiment_dict.json")
    try:
        if os.path.exists(dict_path):
            with open(dict_path, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception:
        pass
    # 内置词典（简化版）
    default_dict = {
        "降息": 0.8, "宽松": 0.6, "刺激": 0.7, "扶持": 0.5, "减税": 0.6, "降准": 0.7,
        "放水": 0.5, "托底": 0.4, "纾困": 0.5, "稳增长": 0.6, "利好": 0.6, "超预期": 0.7,
        "大涨": 0.8, "暴涨": 0.9, "拉升": 0.6, "反弹": 0.5, "普涨": 0.6, "领涨": 0.7,
        "突破": 0.6, "新高": 0.7, "涨停": 0.9, "牛市": 0.8, "上扬": 0.5, "攀升": 0.5,
        "走强": 0.5, "回暖": 0.4, "复苏": 0.5, "修复": 0.4, "企稳": 0.3, "筑底": 0.3,
        "放量": 0.5, "增量": 0.4, "增长": 0.5, "盈利": 0.6, "扭亏": 0.7, "超盈": 0.7,
        "预喜": 0.6, "分红": 0.5, "回购": 0.5, "增持": 0.5, "扩产": 0.4, "订单": 0.4,
        "营收": 0.4, "净利润": 0.5, "毛利率": 0.4, "市占率": 0.4, "龙头": 0.4,
        "独角兽": 0.5, "技术突破": 0.7, "国产替代": 0.6, "自主可控": 0.6, "商业化": 0.5,
        "量产": 0.5, "迭代": 0.4, "升级": 0.4, "转型": 0.4, "出海": 0.4, "全球化": 0.4,
        "净流入": 0.5, "加仓": 0.5, "抄底": 0.4, "布局": 0.4, "配置": 0.3,
        "北上资金": 0.5, "南向资金": 0.5, "ETF申购": 0.4, "融资盘": 0.3,
        "信心": 0.4, "乐观": 0.5, "看多": 0.6, "评级上调": 0.6, "目标价上调": 0.6,
        "增持评级": 0.5, "买入评级": 0.6, "强烈推荐": 0.7, "机构调研": 0.4, "举牌": 0.5,
        "战投": 0.5, "定增": 0.4, "重组": 0.4, "借壳": 0.5, "人工智能": 0.4, "大模型": 0.5,
        "芯片": 0.4, "半导体": 0.4, "新能源": 0.4, "光伏": 0.4, "储能": 0.4, "氢能": 0.4,
        "低空经济": 0.5, "商业航天": 0.5, "创新药": 0.5, "生物医药": 0.4, "数字经济": 0.4,
        "数据要素": 0.4, "算力": 0.5, "5G": 0.3, "6G": 0.4, "物联网": 0.3, "机器人": 0.4,
        "自动驾驶": 0.4, "特斯拉": 0.3, "英伟达": 0.4, "苹果": 0.2, "微软": 0.2,
        "美联储": 0.1, "欧央行": 0.1, "央行": 0.1, "IMF": 0.1, "世行": 0.1, "G20": 0.1,
        "APEC": 0.1, "金砖": 0.1, "RCEP": 0.2, "原油": 0.1, "黄金": 0.1, "铜": 0.1,
        "铁矿石": 0.1, "农产品": 0.1, "大宗商品": 0.1, "贸易协议": 0.5, "停火": 0.6,
        "加息": -0.7, "紧缩": -0.6, "收紧": -0.5, "监管": -0.3, "打压": -0.5,
        "加税": -0.6, "提准": -0.7, "去杠杆": -0.5, "严控": -0.4, "衰退": -0.8,
        "利空": -0.6, "不及预期": -0.6, "低于预期": -0.6, "悲观": -0.5, "滞胀": -0.7,
        "通缩": -0.6, "赤字": -0.4, "债务危机": -0.9, "违约": -0.9, "暴雷": -0.9,
        "黑天鹅": -0.8, "灰犀牛": -0.7, "系统性风险": -0.8, "金融危机": -1.0,
        "经济下行": -0.6, "硬着陆": -0.8, "大跌": -0.8, "暴跌": -0.9, "跳水": -0.7,
        "下挫": -0.6, "普跌": -0.6, "领跌": -0.7, "破位": -0.6, "新低": -0.7,
        "跌停": -0.9, "熊市": -0.8, "阴跌": -0.5, "回调": -0.3, "走弱": -0.5,
        "萎靡": -0.5, "低迷": -0.5, "震荡": -0.2, "缩量": -0.4, "地量": -0.4,
        "抛压": -0.5, "恐慌": -0.7, "下滑": -0.5, "亏损": -0.6, "预亏": -0.7,
        "预减": -0.5, "财务造假": -1.0, "信披违规": -0.8, "立案调查": -0.9,
        "监管函": -0.7, "问询函": -0.6, "警示函": -0.7, "退市": -1.0, "ST": -0.8,
        "*ST": -0.9, "停牌": -0.4, "重组失败": -0.7, "商誉减值": -0.7, "资产减值": -0.6,
        "计提": -0.5, "债务逾期": -0.8, "流动性危机": -0.9, "资金链断裂": -1.0,
        "裁员": -0.5, "关店": -0.4, "减产": -0.4, "停产": -0.5, "召回": -0.6, "诉讼": -0.5,
        "制裁": -0.7, "实体清单": -0.8, "净流出": -0.5, "减仓": -0.5, "清仓": -0.7,
        "撤离": -0.6, "抛售": -0.6, "北上资金流出": -0.5, "南向资金流出": -0.5,
        "ETF赎回": -0.4, "爆仓": -0.9, "强平": -0.8, "恐慌指数": -0.7, "避险": -0.4,
        "看空": -0.6, "评级下调": -0.6, "目标价下调": -0.6, "减持评级": -0.5,
        "卖出评级": -0.6, "机构出逃": -0.6, "股东减持": -0.5, "解禁": -0.4,
        "套现": -0.5, "质押": -0.4, "平仓线": -0.7, "技术封锁": -0.7, "卡脖子": -0.8,
        "断供": -0.7, "专利战": -0.5, "反垄断": -0.5, "天价罚单": -0.7, "安全审查": -0.5,
        "数据泄露": -0.7, "隐私": -0.5, "缺陷": -0.6, "事故": -0.7, "爆炸": -0.9,
        "火灾": -0.8, "罢工": -0.6, "贸易战": -0.7, "关税": -0.5, "脱钩": -0.6,
        "地缘冲突": -0.6, "战争": -1.0, "暴乱": -0.8, "政变": -0.9, "恐袭": -0.9,
        "疫情": -0.7, "封锁": -0.6, "隔离": -0.5, "供应链中断": -0.7, "能源危机": -0.7,
        "粮食危机": -0.8
    }
    # 写入文件供后续使用
    try:
        with open(dict_path, 'w', encoding='utf-8') as f:
            json.dump(default_dict, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    return default_dict

# 加载情感词典
SENTIMENT_DICT = load_sentiment_dict()

class ToastNotification(QFrame):
    def __init__(self, parent, message, duration=2000, bg_color="#10B981"):
        super().__init__(parent)
        self.setObjectName("Toast")
        self.setStyleSheet(f"""
            #Toast {{
                background-color: {bg_color};
                border-radius: 8px;
                padding: 10px 20px;
            }}
        """)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        label = QLabel(message)
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        label.setStyleSheet("color: #1E293B; font-size: 13px; font-weight: 500;")
        layout.addWidget(label)
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint | Qt.WindowType.ToolTip)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        self.setAttribute(Qt.WidgetAttribute.WA_ShowWithoutActivating)
        if parent:
            parent_geo = parent.geometry()
            self.adjustSize()
            x = parent_geo.x() + (parent_geo.width() - self.width()) // 2
            y = parent_geo.y() + parent_geo.height() - self.height() - 50
            self.move(x, y)
        self.opacity_anim = QPropertyAnimation(self, b"windowOpacity")
        self.opacity_anim.setDuration(200)
        self.opacity_anim.setStartValue(0.0)
        self.opacity_anim.setEndValue(1.0)
        self.opacity_anim.start()
        QTimer.singleShot(duration, self.fade_out)

    def fade_out(self):
        self.opacity_anim = QPropertyAnimation(self, b"windowOpacity")
        self.opacity_anim.setDuration(300)
        self.opacity_anim.setStartValue(1.0)
        self.opacity_anim.setEndValue(0.0)
        self.opacity_anim.finished.connect(self.deleteLater)
        self.opacity_anim.start()


class DragDropFrame(QFrame):
    def __init__(self, title, file_type_name, parent=None):
        super().__init__(parent)
        self.title = title
        self.file_type_name = file_type_name
        self.file_paths = []
        self.setAcceptDrops(True)
        self.setMinimumHeight(100)
        self.setMaximumHeight(150)
        self.setStyleSheet("""
            DragDropFrame {
                border: 2px dashed #CBD5E1;
                border-radius: 12px;
                background-color: #F8FAFC;
            }
            DragDropFrame:hover {
                border-color: #1E5F74;
                background-color: #F1F5F9;
            }
        """)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(8)

        title_layout = QHBoxLayout()
        icon_label = QLabel("📂")
        icon_label.setStyleSheet("font-size: 16px;")
        title_label = QLabel(self.title)
        title_label.setStyleSheet("font-weight: bold; color: #1E293B; font-size: 14px;")
        title_layout.addWidget(icon_label)
        title_layout.addWidget(title_label)
        title_layout.addStretch()

        self.browse_btn = QPushButton("浏览文件")
        self.browse_btn.setStyleSheet("""
            QPushButton {
                background-color: transparent;
                border: 1px solid #CBD5E1;
                border-radius: 6px;
                padding: 4px 12px;
                font-size: 12px;
                color: #475569;
            }
            QPushButton:hover {
                background-color: #FFFFFF;
                border-color: #1E5F74;
            }
        """)
        self.browse_btn.clicked.connect(self.browse_files)
        title_layout.addWidget(self.browse_btn)
        layout.addLayout(title_layout)

        hint_label = QLabel(f"拖拽{self.file_type_name}文件至此，或点击「浏览文件」")
        hint_label.setStyleSheet("color: #64748B; font-size: 12px;")
        hint_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(hint_label)

        self.file_list = QListWidget()
        self.file_list.setStyleSheet("""
            QListWidget {
                border: none;
                background: transparent;
                font-size: 12px;
                color: #1E293B;
            }
            QListWidget::item {
                padding: 4px 0;
            }
        """)
        self.file_list.setMaximumHeight(50)
        layout.addWidget(self.file_list)

    def browse_files(self):
        file_paths, _ = QFileDialog.getOpenFileNames(
            self, f"选择{self.file_type_name}文件", "",
            "数据文件 (*.csv *.xlsx *.xls);;所有文件 (*.*)"
        )
        if file_paths:
            self.add_files(file_paths)

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self.setStyleSheet("""
                DragDropFrame {
                    border: 2px solid #1E5F74;
                    border-radius: 12px;
                    background-color: #E6F0F3;
                }
            """)
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self.setStyleSheet("""
            DragDropFrame {
                border: 2px dashed #CBD5E1;
                border-radius: 12px;
                background-color: #F8FAFC;
            }
            DragDropFrame:hover {
                border-color: #1E5F74;
                background-color: #F1F5F9;
            }
        """)
        event.accept()

    def dropEvent(self, event: QDropEvent):
        self.setStyleSheet("""
            DragDropFrame {
                border: 2px dashed #CBD5E1;
                border-radius: 12px;
                background-color: #F8FAFC;
            }
            DragDropFrame:hover {
                border-color: #1E5F74;
                background-color: #F1F5F9;
            }
        """)
        files = [u.toLocalFile() for u in event.mimeData().urls()]
        self.add_files(files)
        event.acceptProposedAction()

    def add_files(self, file_paths):
        valid_exts = ('.csv', '.xlsx', '.xls')
        new_files = []
        for fp in file_paths:
            ext = os.path.splitext(fp)[1].lower()
            if ext in valid_exts:
                if fp not in self.file_paths:
                    self.file_paths.append(fp)
                    self.file_list.addItem(f"✅ {os.path.basename(fp)}")
                    new_files.append(fp)
            else:
                self.file_list.addItem(f"❌ {os.path.basename(fp)} (格式不支持)")

        if new_files:
            if hasattr(self.parent(), 'on_files_added'):
                self.parent().on_files_added(new_files)

    def clear_files(self):
        self.file_paths.clear()
        self.file_list.clear()

    def get_file_paths(self):
        return self.file_paths.copy()


class PersonalizedAnalysisWorker(QThread):
    finished = pyqtSignal(str, str, dict, str, object)
    error = pyqtSignal(str)
    toast = pyqtSignal(str, bool)

    def __init__(self, stock_name, user_category, future_plan, market_sentiment, hot_news):
        super().__init__()
        self.stock_name = stock_name
        self.user_category = user_category
        self.future_plan = future_plan
        self.market_sentiment = market_sentiment
        self.hot_news = hot_news
        self._is_running = True

    def run(self):
        try:
            if not self._is_running: return
            self.toast.emit("正在分析股票数据...", True)
            
            # AI分析股票历史和分类
            ai_category = self._analyze_stock_category()
            if not self._is_running: return

            # 获取股票历史数据
            stock_data = self._get_stock_data()
            if not self._is_running: return

            # 分析热点情感影响
            sentiment_analysis, keyword_scores = self._analyze_sentiment_impact()
            if not self._is_running: return

            # 生成股票预测（传入热点新闻）
            stock_prediction = self._generate_stock_prediction()
            if not self._is_running: return

            # 生成分析报告
            report = self._generate_report(ai_category, sentiment_analysis, keyword_scores, stock_data)
            
            self.finished.emit(report, ai_category, keyword_scores, stock_prediction, stock_data)
        except Exception as e:
            if self._is_running:
                self.error.emit(str(e))

    def stop(self):
        self._is_running = False
        self.wait()

    def _analyze_stock_category(self):
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        prompt = f"""
        请分析股票名称"{self.stock_name}"的业务模式，从以下选项中选择最匹配的一个分类：
        [科技股, 外贸股, 消费股, 金融股, 新能源股, 医药股, 其他]
        只需返回分类名称，不要任何解释。例如：科技股
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        raw = response.choices[0].message.content.strip()
        # 提取分类词
        possible = ["科技股", "外贸股", "消费股", "金融股", "新能源股", "医药股", "其他"]
        for cat in possible:
            if cat in raw:
                return cat
        return "其他"
    
    def _get_stock_data(self):
        """获取股票历史数据"""
        if not AKSHARE_AVAILABLE:
            return None
        
        try:
            # 尝试获取股票代码
            stock_code = self._get_stock_code()
            if not stock_code:
                return None
            
            # 获取日K线数据
            end_date = datetime.now().strftime("%Y%m%d")
            start_date = (datetime.now() - timedelta(days=365)).strftime("%Y%m%d")
            
            # 使用akshare获取股票数据
            stock_data = ak.stock_zh_a_hist(symbol=stock_code, start_date=start_date, end_date=end_date, adjust="qfq")
            return stock_data
        except Exception:
            return None
    
    def _get_stock_code(self):
        """根据股票名称获取股票代码"""
        try:
            # 使用akshare的股票名称查询功能
            stock_info = ak.stock_zh_a_spot_em()
            stock_info = stock_info[stock_info['名称'] == self.stock_name]
            if not stock_info.empty:
                return stock_info.iloc[0]['代码']
            return None
        except Exception:
            return None

    def _analyze_sentiment_impact(self):
        # 分析热点对不同类型股票的影响
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        news_summary = "\n".join([f"• {item}" for item in self.hot_news[:5]])
        prompt = f"""
        基于以下本周市场热点，分析股票"{self.stock_name}"可能受到的影响：
        {news_summary}
        用户未来计划：{self.future_plan if self.future_plan else '无'}
        当前市场情绪指数：{self.market_sentiment:.3f}
        请用一段话简要说明主要热点对该股票的情感影响，并给出短期和长期展望。
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        sentiment_analysis = response.choices[0].message.content
        
        # 让AI生成关键词情感得分
        keyword_scores = self._generate_keyword_scores()
        
        return sentiment_analysis, keyword_scores
    
    def _generate_keyword_scores(self):
        # 让AI为热点新闻中的关键词生成情感得分
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        hot_news_text = "\n".join(self.hot_news)
        prompt = f"""
        请分析以下热点新闻，并为其中的关键词生成情感得分：
        热点新闻：
        {hot_news_text}
        
        股票：{self.stock_name}
        
        请为新闻中的重要关键词生成情感得分，得分范围为-1.0到1.0，其中正值表示积极影响，负值表示消极影响。
        格式要求：
        关键词1: 得分1
        关键词2: 得分2
        ...
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        
        keyword_scores = {}
        try:
            for line in response.choices[0].message.content.split('\n'):
                line = line.strip()
                if ':' in line:
                    parts = line.split(':', 1)
                    if len(parts) == 2:
                        keyword = parts[0].strip()
                        score_str = parts[1].strip()
                        try:
                            score = float(score_str)
                            keyword_scores[keyword] = score
                        except ValueError:
                            pass
        except Exception:
            pass
        
        return keyword_scores

    def _generate_stock_prediction(self):
        # 生成股票预测，结合热点
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        news_summary = "\n".join([f"• {item}" for item in self.hot_news[:3]])
        prompt = f"""
        基于以下本周市场热点和情绪，预测股票"{self.stock_name}"的短期（1-3个月）表现：
        {news_summary}
        当前市场情绪指数：{self.market_sentiment:.3f}
        请给出明确的预测结果（上涨/下跌/震荡）并简要说明关键影响因素，字数在100字以内。
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        return response.choices[0].message.content

    def _generate_report(self, ai_category, sentiment_analysis, keyword_scores, stock_data):
        # 用户分类部分
        user_cat_text = self.user_category if self.user_category != "请选择" else "未指定"
        # 关键词部分（富文本）
        keyword_lines = []
        if keyword_scores:
            sorted_keywords = sorted(keyword_scores.items(), key=lambda x: abs(x[1]), reverse=True)[:10]
            for word, score in sorted_keywords:
                color = "#2E7D32" if score > 0 else "#D32F2F"
                keyword_lines.append(f"<li><span style='color:{color};'>{word}</span>: {score:.2f}</li>")
            keyword_section = "<ul>" + "".join(keyword_lines) + "</ul>"
        else:
            keyword_section = "<p>无相关关键词</p>"
        
        # 股票历史数据分析
        stock_data_analysis = ""
        if stock_data is not None:
            # 简单分析历史数据
            recent_close = stock_data['收盘'].iloc[-1]
            avg_close = stock_data['收盘'].mean()
            max_close = stock_data['收盘'].max()
            min_close = stock_data['收盘'].min()
            
            stock_data_analysis = f"<p>近一年股价：最高 {max_close:.2f}，最低 {min_close:.2f}，平均 {avg_close:.2f}，最新 {recent_close:.2f}</p>"
        else:
            stock_data_analysis = "<p>无法获取股票历史数据</p>"
        
        # 生成未来趋势预测和投资建议
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        prediction_prompt = f"""
        基于以下信息，为股票"{self.stock_name}"生成简洁的未来趋势预测和投资建议：
        股票分类：{ai_category}
        热点影响分析：{sentiment_analysis}
        股票历史数据：{stock_data_analysis.replace('<p>', '').replace('</p>', '')}
        请提供：
        1. 短期（1-3个月）趋势预测（一句话总结）
        2. 长期（6-12个月）趋势预测（一句话总结）
        3. 具体的投资建议（2-3点）
        要求：简洁明了，直接给出结论，不要任何引言或开场白。
        """
        prediction_response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prediction_prompt}],
            stream=False
        )
        prediction_content = prediction_response.choices[0].message.content
        
        return f"""
        <h2>📊 {self.stock_name} 智能分析报告</h2>
        <h3>🤖 AI 智能分类</h3>
        <p><b>{ai_category}</b> （基于公开数据与业务模式）</p>
        <h3>👤 用户自定义定位</h3>
        <p>{user_cat_text}</p>
        <h3>🔥 热点影响分析</h3>
        <p>{sentiment_analysis}</p>
        <h3>📊 历史数据概览</h3>
        {stock_data_analysis}
        <h3>🔑 关键词情感赋值</h3>
        {keyword_section}
        <h3>📈 未来趋势预测</h3>
        <p>{prediction_content}</p>
        """

class HotNewsWorker(QThread):
    finished = pyqtSignal(list, dict)
    error = pyqtSignal(str)
    toast = pyqtSignal(str, bool)

    def __init__(self):
        super().__init__()
        self._is_running = True

    def run(self):
        try:
            if not self._is_running: return
            self.toast.emit("正在分析本周市场热点新闻...", True)
            news = self._fetch_weekly_news()
            if not self._is_running: return
            sentiment_score, top_words = self._analyze(news)
            self.finished.emit(news, {
                "score": sentiment_score,
                "top_positive": top_words["positive"],
                "top_negative": top_words["negative"]
            })
        except Exception as e:
            if self._is_running:
                self.error.emit(str(e))

    def stop(self):
        self._is_running = False
        self.wait()

    def _fetch_weekly_news(self):
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        now = datetime.now()
        start_of_week = (now - timedelta(days=now.weekday())).strftime("%m月%d日")
        end_of_week = now.strftime("%m月%d日")
        prompt = f"""
        请扮演专业的金融市场分析师。从财新网、第一财经、证券时报、中国证券报、财联社、中国人民银行、证监会、上交所、深交所、路透社、彭博社、金融时报、华尔街日报等来源，生成8-10条本周（{start_of_week}至{end_of_week}）最重要的金融/融资市场快讯。
        每条快讯一句话，简洁有力，包含具体来源。
        格式如下：
        • [来源] 热点内容
        • [来源] 热点内容
        • [来源] 热点内容
        ...
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False,
            extra_body={"enable_thinking": True}
        )
        answer = response.choices[0].message.content
        news = []
        for line in answer.split('\n'):
            line = line.strip()
            if line.startswith('•'):
                text = line.lstrip('•').strip()
                if text:
                    news.append(text)
        if not news:
            raise ValueError("AI未返回有效新闻")
        return news[:10]

    def _analyze(self, news_list):
        # 使用AI分析新闻情感
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        news_text = "\n".join(news_list)
        prompt = f"""
        请分析以下市场热点新闻的整体情感倾向，并提取出最积极和最消极的关键词：
        新闻：
        {news_text}
        
        请返回：
        1. 一个整体情感得分（-1.0到1.0，正值表示积极，负值表示消极）
        2. 最积极的3个关键词及其得分
        3. 最消极的3个关键词及其得分
        
        格式要求：
        情感得分: X.XX
        积极关键词:
        关键词1: 得分1
        关键词2: 得分2
        关键词3: 得分3
        消极关键词:
        关键词1: 得分1
        关键词2: 得分2
        关键词3: 得分3
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        
        # 解析AI响应
        content = response.choices[0].message.content
        sentiment_score = 0.0
        positive = []
        negative = []
        
        try:
            lines = content.split('\n')
            for i, line in enumerate(lines):
                line = line.strip()
                if line.startswith('情感得分:'):
                    try:
                        sentiment_score = float(line.split(':', 1)[1].strip())
                    except ValueError:
                        pass
                elif line.startswith('积极关键词:'):
                    # 提取接下来的3行
                    for j in range(1, 4):
                        if i+j < len(lines):
                            kw_line = lines[i+j].strip()
                            if ':' in kw_line:
                                parts = kw_line.split(':', 1)
                                if len(parts) == 2:
                                    kw = parts[0].strip()
                                    try:
                                        score = float(parts[1].strip())
                                        positive.append((kw, score))
                                    except ValueError:
                                        pass
                elif line.startswith('消极关键词:'):
                    # 提取接下来的3行
                    for j in range(1, 4):
                        if i+j < len(lines):
                            kw_line = lines[i+j].strip()
                            if ':' in kw_line:
                                parts = kw_line.split(':', 1)
                                if len(parts) == 2:
                                    kw = parts[0].strip()
                                    try:
                                        score = float(parts[1].strip())
                                        negative.append((kw, score))
                                    except ValueError:
                                        pass
        except Exception:
            pass
        
        return sentiment_score, {"positive": positive, "negative": negative}


# 主界面 
class Module4Interface(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("predict_interface")
        self.uploaded_files = []
        self.market_sentiment = 0.0
        self.hot_news = []
        self.news_worker = None
        self.analysis_worker = None
        self.stock_prediction_label = None
        self.init_ui()
        self.load_hot_news()

    def init_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(24, 24, 24, 24)
        main_layout.setSpacing(20)

        title_label = QLabel("📈 智能融资标的分析")
        title_label.setStyleSheet("font-size:20px; font-weight:bold; color:#0A2647;")
        main_layout.addWidget(title_label)

        #  股票智能预测卡片
        analysis_card = QFrame()
        analysis_card.setObjectName("AnalysisCard")
        analysis_card.setStyleSheet("""
            #AnalysisCard {
                background-color: #FFFFFF;
                border-radius: 16px;
                border: 1px solid #E2E8F0;
            }
        """)
        analysis_layout = QVBoxLayout(analysis_card)
        analysis_layout.setContentsMargins(20, 20, 20, 20)
        analysis_layout.setSpacing(16)

        analysis_header = QHBoxLayout()
        analysis_header.addWidget(QLabel("📊 股票智能预测"))
        analysis_layout.addLayout(analysis_header)

        # 股票名称输入
        stock_layout = QHBoxLayout()
        stock_layout.addWidget(QLabel("股票名称:"))
        self.stock_name_edit = QLineEdit()
        self.stock_name_edit.setPlaceholderText("输入股票名称，例如：贵州茅台、阿里巴巴")
        self.stock_name_edit.setStyleSheet("padding: 8px; border: 1px solid #E2E8F0; border-radius: 6px; font-size:14px;")
        self.stock_name_edit.setMinimumWidth(400)
        stock_layout.addWidget(self.stock_name_edit)
        analysis_layout.addLayout(stock_layout)

        # AI分类结果显示
        self.ai_category_label = QLabel("🤖 AI分类结果: 未分析")
        self.ai_category_label.setStyleSheet("font-size:14px; color:#0A2647; font-weight:500;")
        analysis_layout.addWidget(self.ai_category_label)

        # 用户自定义定位
        user_layout = QHBoxLayout()
        user_layout.addWidget(QLabel("自定义定位 (可选):"))
        self.user_category_combo = QComboBox()
        self.user_category_combo.addItems(["请选择", "科技股", "外贸股", "消费股", "金融股", "新能源股", "医药股", "其他"])
        self.user_category_combo.setEditable(True)
        self.user_category_combo.setStyleSheet("padding: 6px; border: 1px solid #E2E8F0; border-radius: 6px; font-size:14px;")
        user_layout.addWidget(self.user_category_combo)
        analysis_layout.addLayout(user_layout)

        # 未来计划
        plan_label = QLabel("未来计划 (可选):")
        analysis_layout.addWidget(plan_label)
        self.plan_edit = QTextEdit()
        self.plan_edit.setPlaceholderText("例如：国内服装股未来5年要发展国外市场")
        self.plan_edit.setMaximumHeight(100)
        self.plan_edit.setStyleSheet("padding: 8px; border: 1px solid #E2E8F0; border-radius: 6px; font-size:14px;")
        analysis_layout.addWidget(self.plan_edit)

        button_layout = QHBoxLayout()
        self.analyze_btn = QPushButton("🔍 开始智能分析")
        self.analyze_btn.clicked.connect(self.start_analysis)
        self.analyze_btn.setStyleSheet("""
            QPushButton {
                background-color: #0A2647;
                border: none;
                border-radius: 8px;
                padding: 12px 24px;
                font-size: 14px;
                color: white;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #1E5F74;
            }
        """)
        button_layout.addWidget(self.analyze_btn)
        
        self.export_btn = QPushButton("📤 导出报告")
        self.export_btn.clicked.connect(self.export_report)
        self.export_btn.setStyleSheet("""
            QPushButton {
                background-color: #FFFFFF;
                border: 1px solid #0A2647;
                border-radius: 8px;
                padding: 12px 24px;
                font-size: 14px;
                color: #0A2647;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #F8FAFC;
            }
        """)
        button_layout.addWidget(self.export_btn)
        button_layout.addStretch()
        analysis_layout.addLayout(button_layout)

        main_layout.addWidget(analysis_card, 1)

        # 市场热点情绪与分析报告合并卡片
        combined_card = QFrame()
        combined_card.setObjectName("CombinedCard")
        combined_card.setStyleSheet("""
    #CombinedCard {
        background-color: #FFFFFF;
        border-radius: 16px;
        border: 1px solid #E2E8F0;
    }
""")
        combined_layout = QVBoxLayout(combined_card)
        combined_layout.setContentsMargins(20, 20, 20, 20)
        combined_layout.setSpacing(16)

        # 创建一个滚动区域，包含市场热点情绪和分析报告
        self.combined_scroll = QScrollArea()
        self.combined_scroll.setWidgetResizable(True)
        self.combined_scroll.setStyleSheet("border: none; background-color: transparent;")
        
        # 滚动区域内容
        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll_layout.setContentsMargins(0, 0, 0, 0)
        scroll_layout.setSpacing(16)

        # 市场热点情绪部分
        sentiment_section = QVBoxLayout()
        sentiment_section.setSpacing(8)
        
        sentiment_header = QHBoxLayout()
        sentiment_header.addWidget(QLabel("🔥 本周市场热点情绪"))
        self.refresh_sentiment_btn = QPushButton("刷新")
        self.refresh_sentiment_btn.clicked.connect(self.load_hot_news)
        self.refresh_sentiment_btn.setStyleSheet("""
            QPushButton {
                background-color: transparent;
                border: 1px solid #D0D7DE;
                border-radius: 6px;
                padding: 4px 12px;
                font-size: 12px;
            }
            QPushButton:hover {
                background-color: #F8FAFC;
                border-color: #1E5F74;
            }
        """)
        sentiment_header.addStretch()
        sentiment_header.addWidget(self.refresh_sentiment_btn)
        sentiment_section.addLayout(sentiment_header)

        self.news_label = QLabel("正在获取热点新闻...")
        self.news_label.setWordWrap(True)
        self.news_label.setStyleSheet("font-size:14px; color:#1E293B; line-height:1.6;")
        sentiment_section.addWidget(self.news_label)

        self.keyword_analysis_label = QLabel("关键词分析: 未分析")
        self.keyword_analysis_label.setWordWrap(True)
        self.keyword_analysis_label.setStyleSheet("font-size:14px; color:#64748B; margin-top:8px; line-height:1.6;")
        sentiment_section.addWidget(self.keyword_analysis_label)

        # 股票预测部分
        self.stock_prediction_label = QLabel("股票预测: 请输入股票名称并分析")
        self.stock_prediction_label.setWordWrap(True)
        self.stock_prediction_label.setStyleSheet("font-size:14px; color:#1E5F74; font-weight:500; margin-top:8px; line-height:1.6;")
        sentiment_section.addWidget(self.stock_prediction_label)

        scroll_layout.addLayout(sentiment_section)

        # 分析报告部分
        report_section = QVBoxLayout()
        report_section.setSpacing(8)
        
        report_header = QHBoxLayout()
        report_header.addWidget(QLabel("📋 分析报告"))
        report_section.addLayout(report_header)

        self.result_content = QLabel()
        self.result_content.setWordWrap(True)
        self.result_content.setTextFormat(Qt.TextFormat.RichText)  # 启用富文本
        self.result_content.setStyleSheet("font-size:14px; color:#1E293B; line-height:1.8;")
        self.result_content.setAlignment(Qt.AlignmentFlag.AlignTop)
        report_section.addWidget(self.result_content)

        scroll_layout.addLayout(report_section)

        # 设置滚动区域内容
        self.combined_scroll.setWidget(scroll_content)
        combined_layout.addWidget(self.combined_scroll)

        main_layout.addWidget(combined_card, 10)

    def show_toast(self, message, is_success=True):
        bg_color = "#10B981" if is_success else "#EF4444"
        toast = ToastNotification(self.window(), message, duration=2000, bg_color=bg_color)
        toast.show()

    def load_hot_news(self):
        self.refresh_sentiment_btn.setEnabled(False)
        # 安全停止之前的线程
        if self.news_worker and self.news_worker.isRunning():
            self.news_worker.stop()
        self.news_worker = HotNewsWorker()
        self.news_worker.toast.connect(self.show_toast)
        self.news_worker.finished.connect(self.on_news_ready)
        self.news_worker.error.connect(self.on_news_error)
        self.news_worker.start()

    def on_news_ready(self, news, sentiment_data):
        self.refresh_sentiment_btn.setEnabled(True)
        self.hot_news = news
        self.market_sentiment = sentiment_data["score"]
        
        # 显示热点新闻
        news_text = "\n".join([f"• {item}" for item in news])
        self.news_label.setText(news_text)
        
        # 隐藏关键词分析
        self.keyword_analysis_label.setText("")

    def on_news_error(self, err):
        self.refresh_sentiment_btn.setEnabled(True)
        self.news_label.setText(f"❌ 获取热点新闻失败: {err}")
        self.keyword_analysis_label.setText("关键词分析: 获取失败")

    def start_analysis(self):
        stock_name = self.stock_name_edit.text().strip()
        if not stock_name:
            self.show_toast("请输入股票名称", False)
            return
        user_category = self.user_category_combo.currentText()
        future_plan = self.plan_edit.toPlainText().strip()
        self.analyze_btn.setEnabled(False)
        self.result_content.setText("🤖 AI 正在分析中，请稍候...")
        # 安全停止之前的线程
        if self.analysis_worker and self.analysis_worker.isRunning():
            self.analysis_worker.stop()
        self.analysis_worker = PersonalizedAnalysisWorker(
            stock_name, user_category, future_plan, self.market_sentiment, self.hot_news
        )
        self.analysis_worker.toast.connect(self.show_toast)
        self.analysis_worker.finished.connect(self.on_analysis_ready)
        self.analysis_worker.error.connect(self.on_analysis_error)
        self.analysis_worker.start()

    def on_analysis_ready(self, answer, ai_category, keyword_scores, stock_prediction, stock_data):
        self.analyze_btn.setEnabled(True)
        self.result_content.setText(answer)
        self.ai_category_label.setText(f"🤖 AI分类结果: {ai_category}")
        self.stock_prediction_label.setText(f"股票预测: {stock_prediction}")
        self.show_toast("分析完成", True)

    def on_analysis_error(self, err):
        self.analyze_btn.setEnabled(True)
        self.result_content.setText(f"❌ 分析失败: {err}")
        self.show_toast(f"分析失败: {err}", False)

    def export_report(self):
        report_content = self.result_content.text()
        if not report_content or "正在分析中" in report_content:
            self.show_toast("请先完成分析再导出报告", False)
            return
        
        file_path, selected_filter = QFileDialog.getSaveFileName(
            self, "导出分析报告", "", 
            "文本文件 (*.txt);;PDF文件 (*.pdf);;Word文件 (*.docx)"
        )
        if not file_path:
            return
        
        try:
            # 去除HTML标签的函数
            def remove_html_tags(text):
                import re
                clean = re.compile('<.*?>')
                return re.sub(clean, '', text)
            
            if file_path.endswith('.txt'):
                # 去除HTML标签后写入
                clean_content = remove_html_tags(report_content)
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(clean_content)
            elif file_path.endswith('.docx'):
                from docx import Document
                doc = Document()
                doc.add_heading('股票智能分析报告', 0)
                # 去除HTML标签后处理
                clean_content = remove_html_tags(report_content)
                for line in clean_content.split('\n'):
                    line = line.strip()
                    if line:
                        if '📊' in line and '智能分析报告' in line:
                            doc.add_heading(line, 1)
                        elif '🤖' in line or '👤' in line or '🔥' in line or '🔑' in line or '📈' in line:
                            doc.add_heading(line, 2)
                        elif line.startswith('1.') or line.startswith('2.') or line.startswith('3.'):
                            doc.add_paragraph(line, style='List Number')
                        elif line.startswith('   -'):
                            doc.add_paragraph(line, style='List Bullet')
                        else:
                            doc.add_paragraph(line)
                doc.save(file_path)
            elif file_path.endswith('.pdf'):
                # 使用 reportlab 生成 PDF
                from reportlab.lib.pagesizes import A4
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.lib.units import inch
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                from reportlab.lib.enums import TA_LEFT
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont

                font_path = get_resource_path("fonts/SIMHEI.ttf")
                if os.path.exists(font_path):
                    pdfmetrics.registerFont(TTFont('SimHei', font_path))
                    font_name = 'SimHei'

                doc = SimpleDocTemplate(file_path, pagesize=A4)
                styles = getSampleStyleSheet()
                style_normal = ParagraphStyle(
                    'Normal',
                    parent=styles['Normal'],
                    fontName=font_name,
                    fontSize=10,
                    leading=14,
                    alignment=TA_LEFT
                )
                style_heading1 = ParagraphStyle(
                    'Heading1',
                    parent=styles['Heading1'],
                    fontName=font_name,
                    fontSize=16,
                    spaceAfter=12
                )
                style_heading2 = ParagraphStyle(
                    'Heading2',
                    parent=styles['Heading2'],
                    fontName=font_name,
                    fontSize=14,
                    spaceAfter=8
                )

                story = []
                # 去除HTML标签
                plain_text = remove_html_tags(report_content)
                for para in plain_text.split('\n'):
                    para = para.strip()
                    if para:
                        if '📊' in para and '智能分析报告' in para:
                            story.append(Paragraph(para, style_heading1))
                        elif '🤖' in para or '👤' in para or '🔥' in para or '🔑' in para or '📈' in para:
                            story.append(Paragraph(para, style_heading2))
                        else:
                            story.append(Paragraph(para, style_normal))
                            story.append(Spacer(1, 0.1*inch))
                doc.build(story)
            else:
                self.show_toast("暂不支持此格式", False)
                return
            self.show_toast("报告导出成功", True)
        except Exception as e:
            self.show_toast(f"导出失败: {str(e)}", False)

    def closeEvent(self, event):
        """窗口关闭时安全停止所有线程"""
        self.stop_all_workers()
        event.accept()

    def stop_all_workers(self):
        """停止所有正在运行的线程"""
        if self.news_worker and self.news_worker.isRunning():
            self.news_worker.stop()
        if self.analysis_worker and self.analysis_worker.isRunning():
            self.analysis_worker.stop()

    def __del__(self):
        self.stop_all_workers()


# 周度情绪分析线程（保留，未在主界面使用） 
class WeeklySentimentWorker(QThread):
    finished = pyqtSignal(dict)
    error = pyqtSignal(str)
    toast = pyqtSignal(str, bool)

    def __init__(self):
        super().__init__()
        self._is_running = True

    def run(self):
        try:
            if not self._is_running: return
            self.toast.emit("正在分析本周市场热点情绪...", True)
            news = self._fetch_weekly_news()
            if not self._is_running: return
            sentiment_score, top_words = self._analyze(news)
            prediction = "📈 偏乐观" if sentiment_score > 0.1 else ("📉 偏悲观" if sentiment_score < -0.1 else "⚖️ 中性")
            self.finished.emit({
                "score": sentiment_score,
                "prediction": prediction,
                "top_positive": top_words["positive"],
                "top_negative": top_words["negative"]
            })
        except Exception as e:
            if self._is_running:
                self.error.emit(str(e))

    def stop(self):
        self._is_running = False
        self.wait()

    def _fetch_weekly_news(self):
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        now = datetime.now()
        start_of_week = (now - timedelta(days=now.weekday())).strftime("%m月%d日")
        end_of_week = now.strftime("%m月%d日")
        prompt = f"""
        请扮演专业的金融市场分析师。生成5条本周（{start_of_week}至{end_of_week}）最重要的金融/融资市场快讯。
        每条快讯一句话，简洁有力。
        格式如下：
        • [热点1] xxx
        • [热点2] xxx
        • [热点3] xxx
        • [热点4] xxx
        • [热点5] xxx
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False,
            extra_body={"enable_thinking": True}
        )
        answer = response.choices[0].message.content
        news = []
        for line in answer.split('\n'):
            line = line.strip()
            if line.startswith('•'):
                text = line.lstrip('•').strip()
                if text:
                    news.append(text)
        if not news:
            raise ValueError("AI未返回有效新闻")
        return news[:5]

    def _analyze(self, news_list):
        # 使用AI分析新闻情感
        client = ZhipuAI(api_key=ZHIPU_API_KEY)
        news_text = "\n".join(news_list)
        prompt = f"""
        请分析以下市场热点新闻的整体情感倾向，并提取出最积极和最消极的关键词：
        新闻：
        {news_text}
        
        请返回：
        1. 一个整体情感得分（-1.0到1.0，正值表示积极，负值表示消极）
        2. 最积极的3个关键词及其得分
        3. 最消极的3个关键词及其得分
        
        格式要求：
        情感得分: X.XX
        积极关键词:
        关键词1: 得分1
        关键词2: 得分2
        关键词3: 得分3
        消极关键词:
        关键词1: 得分1
        关键词2: 得分2
        关键词3: 得分3
        """
        response = client.chat.completions.create(
            model=MODULE1_MODEL,
            messages=[{"role": "user", "content": prompt}],
            stream=False
        )
        
        # 解析AI响应
        content = response.choices[0].message.content
        sentiment_score = 0.0
        positive = []
        negative = []
        
        try:
            lines = content.split('\n')
            for i, line in enumerate(lines):
                line = line.strip()
                if line.startswith('情感得分:'):
                    try:
                        sentiment_score = float(line.split(':', 1)[1].strip())
                    except ValueError:
                        pass
                elif line.startswith('积极关键词:'):
                    # 提取接下来的3行
                    for j in range(1, 4):
                        if i+j < len(lines):
                            kw_line = lines[i+j].strip()
                            if ':' in kw_line:
                                parts = kw_line.split(':', 1)
                                if len(parts) == 2:
                                    kw = parts[0].strip()
                                    try:
                                        score = float(parts[1].strip())
                                        positive.append((kw, score))
                                    except ValueError:
                                        pass
                elif line.startswith('消极关键词:'):
                    # 提取接下来的3行
                    for j in range(1, 4):
                        if i+j < len(lines):
                            kw_line = lines[i+j].strip()
                            if ':' in kw_line:
                                parts = kw_line.split(':', 1)
                                if len(parts) == 2:
                                    kw = parts[0].strip()
                                    try:
                                        score = float(parts[1].strip())
                                        negative.append((kw, score))
                                    except ValueError:
                                        pass
        except Exception:
            pass
        
        return sentiment_score, {"positive": positive, "negative": negative}
